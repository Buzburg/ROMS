"""Local PDF, Office and image extraction. Originals are never modified or executed."""
import io
from pathlib import Path
import shutil
import subprocess
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

BINARY_EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.xlsm', '.pptx', '.odt', '.ods', '.odp',
                     '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp', '.bmp'}
MAX_PAGES = 100
MAX_PIXELS = 40_000_000
MAX_EXTRACTED_CHARS = 8_000_000


def _run(command, deadline, **kwargs):
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise ValueError('Document extraction exceeded its 120-second budget')
    try:
        result = subprocess.run(command, capture_output=True, timeout=min(60, remaining), **kwargs)
    except subprocess.TimeoutExpired as exc:
        raise ValueError('Local document extractor timed out') from exc
    if result.returncode:
        raise ValueError('Local extractor failed: ' + result.stderr.decode('utf-8', errors='replace')[-500:])
    return result.stdout


def _ocr(image, deadline):
    from PIL import Image, ImageOps
    executable = shutil.which('tesseract')
    if not executable:
        raise ValueError('Local OCR requires the tesseract executable and language data')
    if image.width * image.height > MAX_PIXELS:
        raise ValueError('Image exceeds the 40-megapixel OCR limit')
    image = ImageOps.exif_transpose(image)
    if image.mode == 'RGBA':
        background = Image.new('RGBA', image.size, 'white')
        image = Image.alpha_composite(background, image)
    image = ImageOps.autocontrast(image.convert('L'))
    image.thumbnail((3200, 3200))
    data = io.BytesIO(); image.save(data, format='PNG')
    output = _run([executable, 'stdin', 'stdout', '-l', 'eng'], deadline, input=data.getvalue())
    return output.decode('utf-8', errors='replace').strip()


def _zip_check(raw):
    archive = zipfile.ZipFile(io.BytesIO(raw))
    if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
        archive.close()
        raise ValueError('Office archive exceeds the expanded-size limit')
    return archive


def extract_document(path, raw):
    from PIL import Image
    deadline = time.monotonic() + 120
    suffix = path.suffix.lower()
    notes, chunks = [], []
    total_chars = 0

    def add(value):
        nonlocal total_chars
        if value.strip():
            chunks.append(value)
            total_chars += len(value)
        if total_chars > MAX_EXTRACTED_CHARS or len(chunks) > 20000:
            raise ValueError('Extracted document exceeds the text/passage limit')
        if time.monotonic() > deadline:
            raise ValueError('Document extraction exceeded its time budget')

    def image_text(data, label):
        try:
            with Image.open(io.BytesIO(data)) as image:
                count = getattr(image, 'n_frames', 1)
                if count > MAX_PAGES:
                    raise ValueError('Image has more than 100 frames')
                for frame in range(count):
                    image.seek(frame)
                    text = _ocr(image, deadline)
                    if text:
                        add(f'{label}, frame {frame+1} [OCR]:\n{text}')
                    else:
                        notes.append(f'{label}, frame {frame+1}: OCR found no readable text')
        except (OSError, ValueError) as exc:
            raise ValueError(f'{label}: {exc}') from exc

    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('Password-protected PDF cannot be read without a password')
        if len(reader.pages) > MAX_PAGES:
            raise ValueError('PDF exceeds the 100-page limit')
        with tempfile.TemporaryDirectory(prefix='roms-pdf-') as directory:
            source = Path(directory) / 'source.pdf'; source.write_bytes(raw)
            for number, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or '').strip()
                if text:
                    add(f'Page {number}:\n{text}')
                # Include image text even on mixed pages with a substantial text layer.
                if not text or len(page.images):
                    executable = shutil.which('pdftoppm')
                    if not executable:
                        raise ValueError('Scanned PDF extraction requires pdftoppm and tesseract')
                    prefix = Path(directory) / 'page'
                    _run([executable, '-f', str(number), '-l', str(number), '-singlefile',
                          '-scale-to', '2400', '-png', str(source), str(prefix)], deadline)
                    with Image.open(str(prefix)+'.png') as image:
                        ocr = _ocr(image, deadline)
                    if ocr and ' '.join(ocr.split()) != ' '.join(text.split()):
                        add(f'Page {number} [OCR]:\n{ocr}')
                        notes.append(f'Page {number} includes OCR text; recognition may contain errors')
                    elif not text and not ocr:
                        notes.append(f'Page {number}: no readable text detected')
    elif suffix in {'.docx', '.xlsx', '.xlsm', '.pptx', '.odt', '.ods', '.odp'}:
        with _zip_check(raw) as archive:
            if suffix == '.docx':
                from docx import Document
                document = Document(io.BytesIO(raw))
                for item in document.iter_inner_content():
                    if hasattr(item, 'rows'):
                        for row in item.rows:
                            add(' | '.join(cell.text for cell in row.cells))
                    else:
                        add(item.text)
                for section in document.sections:
                    for part in (section.header, section.footer):
                        for paragraph in part.paragraphs:
                            add(paragraph.text)
            elif suffix in {'.xlsx', '.xlsm'}:
                from openpyxl import load_workbook
                values = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in values:
                        if sheet.max_row and sheet.max_row > 100000 or sheet.max_column and sheet.max_column > 1000:
                            raise ValueError('Spreadsheet exceeds row/column extraction bounds')
                        for row, expressions in zip(sheet.iter_rows(), formulas[sheet.title].iter_rows()):
                            cells = []
                            for cell, expression in zip(row, expressions):
                                value = cell.value
                                if value is None and expression.data_type == 'f':
                                    value = f'[formula, no cached value] {expression.value}'
                                if value is not None:
                                    cells.append(f'{cell.coordinate}: {value}')
                            if cells:
                                add(f'Sheet {sheet.title} | ' + '; '.join(cells))
                finally:
                    values.close(); formulas.close()
            elif suffix == '.pptx':
                from pptx import Presentation
                presentation = Presentation(io.BytesIO(raw))
                if len(presentation.slides) > MAX_PAGES:
                    raise ValueError('Presentation exceeds the 100-slide limit')
                for number, slide in enumerate(presentation.slides, 1):
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            add(f'Slide {number}: {shape.text}')
                        if shape.has_table:
                            for row in shape.table.rows:
                                add(f'Slide {number}: ' + ' | '.join(cell.text for cell in row.cells))
                    if slide.has_notes_slide:
                        add(f'Slide {number} notes: {slide.notes_slide.notes_text_frame.text}')
            else:
                xml = archive.read('content.xml')
                if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                    raise ValueError('ODF entity declarations are not accepted')
                root = ET.fromstring(xml)
                for element in root.iter():
                    if element.tag.rsplit('}', 1)[-1] in {'p', 'h'}:
                        add(''.join(element.itertext()))
            images = [name for name in archive.namelist()
                      if name.startswith(('word/media/', 'ppt/media/', 'xl/media/', 'Pictures/'))
                      and Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp', '.bmp'}]
            if len(images) > MAX_PAGES:
                raise ValueError('Office document exceeds the 100-image OCR limit')
            for name in images:
                image_text(archive.read(name), f'Embedded image {Path(name).name}')
            if images:
                notes.append('Embedded images processed by local English OCR')
    else:
        image_text(raw, path.name)
        notes.append('Image text extracted with local English OCR; recognition may contain errors')
    if not chunks:
        raise ValueError('No searchable text was recovered from this document')
    return path.name, 'document', chunks, notes
