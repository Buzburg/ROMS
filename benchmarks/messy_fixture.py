"""Synthetic, labeled messy documents; no user data. Used by tests and evaluation."""
from pathlib import Path
import zipfile


def make_documents(root):
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    root = Path(root); root.mkdir(parents=True, exist_ok=True)
    documents = {
        'repair.md': '---\ntitle: [broken\n---\nDamaged deliveries qualify for a free replacement within thirty days.',
        'inventory.csv': 'Item;Quantity;Quantity\nCobalt bearings;27;available;Warehouse Delta\nAmber bolts;9\n',
        'events.ndjson': '{"service":"orion", "port":8421}\n{broken: Nebula emergency reboot token 7359\n',
        'notes.txt': 'Café repair appointments open at noon. Bring your receipt.',
        'unicode.log': '北極設備\r\nAurora coolant maintenance interval is forty days.\x01',
        'help.html': '<h1>Account access</h1><script>hiddenkey</script><p>Forgotten passwords can be reset using the recovery email.</p>',
    }
    for name, text in documents.items():
        encoding = 'cp1252' if name == 'notes.txt' else 'utf-16' if name == 'unicode.log' else 'utf-8'
        (root/name).write_bytes(text.encode(encoding))
    image = Image.new('RGB', (1600, 400), 'white')
    font_path = Path('/usr/share/fonts/liberation/LiberationSans-Regular.ttf')
    font = ImageFont.truetype(str(font_path), 48) if font_path.is_file() else ImageFont.load_default(size=48)
    ImageDraw.Draw(image).text((45, 100), 'Violet laboratory access code 4826', fill='black', font=font)
    image.save(root/'scan.png'); image.save(root/'scan.pdf', 'PDF', resolution=150)
    image.save(root/'frames.tiff', save_all=True, append_images=[image])
    document = Document(); document.add_paragraph('Saturn warranty lasts eighteen months.')
    table = document.add_table(rows=1, cols=2); table.cell(0,0).text='Support phone'; table.cell(0,1).text='555-0137'
    document.add_picture(str(root/'scan.png'))
    document.save(root/'warranty.docx')
    workbook = Workbook(); sheet = workbook.active; sheet.title='Orders'
    sheet.append(['Item', 'Order number']); sheet.append(['Emerald pump', 91234]); sheet['C2']='=1+2'
    workbook.save(root/'orders.xlsx')
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text='Helios evacuation assembly point is the northern courtyard.'
    presentation.save(root/'safety.pptx')
    with zipfile.ZipFile(root/'meeting.odt', 'w') as archive:
        archive.writestr('content.xml', '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><text:p>Lunar planning meeting takes place every Wednesday at nine.</text:p></office:document-content>')
    writer = PdfWriter(); page = writer.add_blank_page(width=612, height=792)
    font_obj = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font_obj)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 16 Tf 50 700 Td (Copper turbine inspection is due in September.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(root/'inspection.pdf')
    return [
        ('free replacement for broken delivery', ['repair.md']),
        ('Cobalt bearings warehouse quantity', ['inventory.csv']),
        ('Nebula emergency reboot token', ['events.ndjson']),
        ('When are cafe repair appointments?', ['notes.txt']),
        ('Aurora coolant maintenance interval', ['unicode.log']),
        ('I forgot my password how do I recover access?', ['help.html']),
        ('Violet laboratory access code', ['scan.png','scan.pdf','frames.tiff','warranty.docx']),
        ('Saturn warranty duration', ['warranty.docx']),
        ('Emerald pump order number', ['orders.xlsx']),
        ('Helios evacuation assembly location', ['safety.pptx']),
        ('Lunar planning meeting schedule', ['meeting.odt']),
        ('Copper turbine inspection due', ['inspection.pdf']),
    ]
